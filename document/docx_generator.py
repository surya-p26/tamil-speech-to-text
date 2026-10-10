"""Well-formatted editable Tamil DOCX generation for the existing API endpoint.

Function signature is intentionally preserved:
    generate_docx(text: str, output_file: str) -> bool
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .document_layout import parse_document_text


FONT_NAME = "Nirmala UI"
NAVY = "17365D"
BLUE = "234F7E"
TEXT = "202938"
LIGHT_BLUE = "EAF1F8"
BORDER = "C9D5E3"


def _set_run_font(run, name: str = FONT_NAME, size: float | None = None,
                  bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "ta-IN")


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _append_page_number(paragraph) -> None:
    run = paragraph.add_run()
    _set_run_font(run, size=9, color="586579")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, value, end):
        run._r.append(item)


def _add_metadata_table(document: Document, content) -> None:
    # Department is already shown directly below the document title.
    items = [
        ("தேதி", content.date),
        ("குறிப்பு எண்", content.reference_number),
        ("தயாரித்தது", content.prepared_by),
    ]
    items = [(label, value) for label, value in items if value]
    if not items:
        return

    table = document.add_table(rows=len(items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.45)

    for index, (label, value) in enumerate(items):
        label_cell, value_cell = table.rows[index].cells
        label_cell.width = Inches(1.35)
        value_cell.width = Inches(5.45)
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(label_cell, LIGHT_BLUE)
        _set_cell_borders(label_cell)
        _set_cell_borders(value_cell)

        label_para = label_cell.paragraphs[0]
        label_para.paragraph_format.space_after = Pt(0)
        label_run = label_para.add_run(label)
        _set_run_font(label_run, size=10, bold=True, color=NAVY)

        value_para = value_cell.paragraphs[0]
        value_para.paragraph_format.space_after = Pt(0)
        value_run = value_para.add_run(value)
        _set_run_font(value_run, size=10, color=TEXT)

    document.add_paragraph().paragraph_format.space_after = Pt(1)


def generate_docx(text: str, output_file: str) -> bool:
    """Generate a clean, editable DOCX with metadata and the true transcript."""
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Cannot generate a DOCX from empty text.")

    content = parse_document_text(text)
    out_path = Path(output_file)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    document.core_properties.title = content.title
    document.core_properties.subject = "Tamil speech-to-text transcription"
    document.core_properties.author = content.prepared_by or "TamilTranscribe"

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    normal_rpr = normal._element.get_or_add_rPr()
    normal_fonts = normal_rpr.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_rpr.insert(0, normal_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{attr}"), FONT_NAME)

    govt = document.add_paragraph()
    govt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    govt.paragraph_format.space_after = Pt(4)
    run = govt.add_run(content.government)
    _set_run_font(run, size=16, bold=True, color=NAVY)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(3)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(content.title)
    _set_run_font(run, size=17, bold=True, color=NAVY)

    if content.department:
        department = document.add_paragraph()
        department.alignment = WD_ALIGN_PARAGRAPH.CENTER
        department.paragraph_format.space_after = Pt(8)
        run = department.add_run(content.department)
        _set_run_font(run, size=10.5, color="586579")

    # Add a bottom border to the title block for a tidy official-document look.
    p_pr = title._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    _add_metadata_table(document, content)

    section_heading = document.add_paragraph()
    section_heading.paragraph_format.space_before = Pt(8)
    section_heading.paragraph_format.space_after = Pt(6)
    heading_run = section_heading.add_run("படியெழுத்தாக்கம்")
    _set_run_font(heading_run, size=12.5, bold=True, color=BLUE)

    transcript = content.transcript.strip()
    if not transcript:
        transcript = "(எழுத்தாக்க உரை இல்லை)"
    for block in transcript.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.widow_control = True
        lines = block.split("\n")
        for line_index, line in enumerate(lines):
            if line_index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            _set_run_font(run, size=11.5, color=TEXT)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(4)
    label_run = footer.add_run("TamilTranscribe  |  பக்கம் ")
    _set_run_font(label_run, size=9, color="586579")
    _append_page_number(footer)

    document.save(str(out_path))
    return out_path.exists() and out_path.stat().st_size > 0
